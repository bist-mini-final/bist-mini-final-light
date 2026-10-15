"""Read the checked-in fictional workbook fixtures and build honest BI demos.

Only the checksum-verified sample set is accepted. Ordinary user workbooks
continue through the normal RAG materialization path.
"""

from __future__ import annotations

import hashlib
import json
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.workbook.workbook import Workbook

from backend.domains.bi.application.snapshot_builder import BiSnapshotBuilder
from backend.domains.bi.domain.catalog import METRIC_CATALOG, SourceMetricDefinition
from backend.domains.bi.domain.extraction_models import BiMetricExtractionResult
from backend.domains.bi.domain.materialization_models import BiDocumentProfile, BiSnapshotBuildInput
from backend.domains.bi.domain.models import (
    AmountScale,
    AvailableObservation,
    BiDashboardSnapshot,
    BiEvidence,
    BiMaterializationRequest,
    BiPeriod,
    JobId,
    MetricId,
    MetricStatus,
    PeriodId,
    PeriodKind,
    UnavailableObservation,
)


def read_sample_manifest(directory: Path) -> dict[str, Any]:
    manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
    if manifest.get("synthetic") is not True or manifest.get("schema_version") != 1:
        raise ValueError("Only the fictional sample manifest is supported")
    return manifest


def _read_observation(
    cell: Cell | MergedCell,
    expected: Any,
    period: BiPeriod,
    company_name: str,
    label: str,
) -> AvailableObservation | UnavailableObservation:
    if not isinstance(cell, Cell):
        raise ValueError("Sample metric values must use unmerged cells")
    value = cell.value
    if value is None:
        if expected is not None:
            raise ValueError(f"Missing sample value: {cell.coordinate}")
        return UnavailableObservation(
            period_id=period.period_id,
            status=MetricStatus.MISSING,
            reason="not_reported_in_fixture",
        )
    if not isinstance(value, (int, float)) or isinstance(value, bool) or expected is None:
        raise ValueError(f"Invalid sample number: {cell.coordinate}")
    number = Decimal(str(value))
    if abs(number - Decimal(str(expected))) > Decimal("0.00000001"):
        raise ValueError(f"Changed sample value: {cell.coordinate}")
    sheet_name = cell.parent.title
    evidence = BiEvidence(
        cell_id=f"{sheet_name}!{cell.coordinate}",
        sheet_name=sheet_name,
        cell_coord=cell.coordinate,
        source_text=f"Company: {company_name} | Metric: {label} | Period: {period.source_label} | Value: {number} | USD millions",
    )
    return AvailableObservation(
        period_id=period.period_id,
        status=MetricStatus.AVAILABLE,
        raw_value=str(value),
        normalized_value=number,
        evidence=(evidence,),
        notes=("Verified fictional workbook cell; no LLM metric extraction",),
    )


def _read_source_results(
    workbook: Workbook,
    manifest: dict[str, Any],
    company: dict[str, Any],
    periods: tuple[BiPeriod, ...],
) -> tuple[BiMetricExtractionResult, ...]:
    extracted: list[BiMetricExtractionResult] = []
    for sheet_name, rows in manifest["sheet_rows"].items():
        sheet = workbook[sheet_name]
        for column, period in enumerate(periods, 2):
            header = sheet.cell(5, column).value
            if not isinstance(header, datetime) or header.date() != period.end_date:
                raise ValueError(f"Incorrect period header: {company['file']}/{sheet_name}")
            expected = company["values"][column - 2]
            for row, (metric_name, label) in enumerate(rows, 6):
                metric_id = MetricId(metric_name)
                definition = METRIC_CATALOG[metric_id]
                if not isinstance(definition, SourceMetricDefinition):
                    raise ValueError("Sample contains a non-source metric")
                if sheet.cell(row, 1).value != label:
                    raise ValueError(f"Incorrect metric label: {company['file']}/{sheet_name}")
                observation = _read_observation(
                    sheet.cell(row, column),
                    expected[metric_name],
                    period,
                    company["name"],
                    label,
                )
                extracted.append(
                    BiMetricExtractionResult(
                        metric_id=metric_id,
                        period_id=period.period_id,
                        value_kind=definition.value_kind,
                        currency="USD",
                        scale=AmountScale.MILLIONS,
                        observation=observation,
                    )
                )
    return tuple(extracted)


def build_sample_snapshot(
    directory: Path,
    manifest: dict[str, Any],
    company: dict[str, Any],
    request: BiMaterializationRequest,
    *,
    now: datetime | None = None,
) -> BiDashboardSnapshot:
    """Validate every actual workbook cell before using the shared BI calculator."""
    filename = company["file"]
    if Path(filename).name != filename:
        raise ValueError("Sample file must stay inside the sample directory")
    path = directory / filename
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if (
        digest != company["sha256"]
        or request.source.workbook_hash != digest
        or request.source.file_name != filename
    ):
        raise ValueError("Sample workbook hash does not match its manifest and index")
    periods = tuple(
        BiPeriod(
            period_id=PeriodId(f"fy-{year}-12-31"),
            kind=PeriodKind.FY,
            label=f"FY{year}",
            source_label=f"{year}-12-31",
            end_date=date(year, 12, 31),
            ordinal=year,
        )
        for year in manifest["years"]
    )
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        extracted = _read_source_results(workbook, manifest, company, periods)
    finally:
        workbook.close()
    profile = BiDocumentProfile(
        periods=periods,
        currency="USD",
        scale=AmountScale.MILLIONS,
        relevant_sheets=tuple(manifest["sheet_rows"]),
    )
    snapshot = BiSnapshotBuilder().build(
        BiSnapshotBuildInput(
            request=request,
            job_id=JobId(f"sample-{company['id']}-{digest[:12]}"),
            profile=profile,
            extracted=tuple(extracted),
            generated_at=now or datetime.now(UTC),
        )
    )
    return snapshot.model_copy(
        update={
            "refresh": snapshot.refresh.model_copy(
                update={
                    "message": "실습용 가상 데이터의 원본 셀에서 직접 계산한 대시보드입니다. LLM 지표 추출 결과가 아닙니다."
                }
            )
        }
    )
